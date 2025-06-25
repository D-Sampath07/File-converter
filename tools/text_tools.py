# tools/text_tools.py
from fpdf import FPDF
from docx import Document

def txt_to_pdf(input_path, output_path):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.set_font("Arial", size=12)

    with open(input_path, 'r', encoding='utf-8') as f:
        for line in f:
            pdf.cell(0, 10, txt=line.strip(), ln=True)

    pdf.output(output_path)

def txt_to_docx(input_path, output_path):
    doc = Document()
    with open(input_path, 'r', encoding='utf-8') as f:
        for line in f:
            doc.add_paragraph(line.strip())
    doc.save(output_path)

def convert(input_path, output_path, target_format):
    target_format = target_format.lower()
    if target_format == "pdf":
        txt_to_pdf(input_path, output_path)
    elif target_format == "docx":
        txt_to_docx(input_path, output_path)
    else:
        raise ValueError(f"Conversion from TXT to {target_format.upper()} not supported.")
