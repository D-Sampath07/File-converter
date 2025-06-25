# ✅ docx_tools.py (works on Render)
from docx import Document
from pptx import Presentation
from pptx.util import Inches, Pt

def docx_to_txt(input_path, output_path):
    doc = Document(input_path)
    text = '\n'.join([para.text for para in doc.paragraphs])
    with open(output_path, 'w', encoding='utf-8') as f:
        f.write(text)

def docx_to_pptx(input_path, output_path):
    doc = Document(input_path)
    prs = Presentation()
    blank_slide_layout = prs.slide_layouts[6]  # Fully blank slide

    for para in doc.paragraphs:
        if para.text.strip():
            slide = prs.slides.add_slide(blank_slide_layout)
            txBox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(5.5))
            tf = txBox.text_frame
            p = tf.paragraphs[0]
            p.text = para.text.strip()
            p.font.size = Pt(20)

    prs.save(output_path)

def convert(input_path, target_format, output_path):
    target_format = target_format.lower()
    if target_format == 'txt':
        docx_to_txt(input_path, output_path)
    elif target_format == 'pptx':
        docx_to_pptx(input_path, output_path)
    else:
        raise ValueError(f"DOCX to {target_format.upper()} is not supported.")
